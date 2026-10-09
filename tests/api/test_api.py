"""The web API end to end, offline: a FakeFetcher stands in for the network."""

import io
import json
import time
import zipfile

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from scrapebot.api import ApiSettings, create_app
from scrapebot.config import RunConfig
from tests.fakes import FakeFetcher

SHOPIFY_HOME = (
    '<html><script>Shopify.currency = {"active":"USD","rate":"1.0"};</script>cdn.shopify.com</html>'
)
FEED = json.dumps(
    {"products": [{"title": "Cher Sweater", "handle": "cher", "variants": [{"price": "139.00"}]}]}
)
READABLE = "<html><body><p>" + "A boutique in Naples. " * 30 + "</p></body></html>"
RESPONSES = {
    "https://monkees.com": (200, SHOPIFY_HOME),
    "https://monkees.com/products.json?limit=250&page=1": (200, FEED),
    "https://boutique.com": (200, READABLE),
    "https://blocked.com": (403, ""),
}
LINKS = "Cek monkees.com, https://boutique.com dan https://blocked.com. IG: https://instagram.com/x"


@pytest.fixture
def client(tmp_path):
    base = RunConfig.model_validate(
        {"output": {"runs_dir": tmp_path / "runs"}, "fetch": {"cache_dir": tmp_path / "cache"}}
    )
    settings = ApiSettings(base=base, uploads_dir=tmp_path / "uploads", poll_seconds=0.01)
    app = create_app(settings, fetcher_factory=lambda: FakeFetcher(RESPONSES))
    with TestClient(app) as c:
        yield c


def wait_until_finished(client, run_id, timeout=10.0):
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        run = client.get(f"/api/runs/{run_id}").json()
        if run["state"] in ("done", "failed", "interrupted", "stopped"):
            return run
        time.sleep(0.02)
    raise AssertionError(f"run {run_id} did not finish")


def start(client, **body):
    payload = {"source": {"text": LINKS}, "writers": ["xlsx", "csv"], **body}
    return client.post("/api/runs", json=payload)


def test_options_list_the_writers_and_defaults(client):
    body = client.get("/api/options").json()
    assert "parquet" in body["writers"]
    assert body["default_writers"] == ["xlsx", "csv"]


def test_preview_counts_links_before_anything_is_fetched(client):
    body = client.post(
        "/api/preview", json={"source": {"text": LINKS + " monkees.com/shop"}}
    ).json()
    assert body["links"] == 5
    assert body["stores"] == 3
    assert body["duplicates"] == 1
    assert body["skipped"] == {"social_only": 1}


def test_main_scenario_run_check_download(client):
    """PRD section 4: paste, run every store, check the data, download."""
    resp = start(client)
    assert resp.status_code == 202
    full_run = wait_until_finished(client, resp.json()["run_id"])
    assert full_run["state"] == "done"
    assert full_run["mode"] == "full"
    assert full_run["stores_done"] == 3
    assert full_run["links_in"] == full_run["processed"] + full_run["skipped"]
    assert full_run["skipped_by_reason"] == {"social_only": 1}
    assert full_run["status_counts"] == {"ok": 1, "no_products": 1, "blocked": 1}
    assert full_run["products"] == 1
    assert {d["key"] for d in full_run["downloads"]} >= {
        "report",
        "summary",
        "xlsx",
        "csv",
        "manifest",
    }

    xlsx = client.get(f"/api/runs/{full_run['run_id']}/download/xlsx")
    assert xlsx.status_code == 200
    sheets = load_workbook(io.BytesIO(xlsx.content), read_only=True).sheetnames
    assert sheets == [
        "runs",
        "inputs",
        "stores",
        "products",
        "pages",
        "contacts",
        "changes",
        "llm_calls",
    ]

    csv_zip = client.get(f"/api/runs/{full_run['run_id']}/download/csv")
    assert csv_zip.status_code == 200
    assert "products.csv" in zipfile.ZipFile(io.BytesIO(csv_zip.content)).namelist()

    listed = client.get("/api/runs").json()
    assert [r["run_id"] for r in listed] == [full_run["run_id"]]


def test_upload_then_run_a_spreadsheet(client):
    csv_bytes = b"store_name,website\nMonkees,https://monkees.com\nNo site,\n"
    up = client.post("/api/uploads", files={"file": ("../../evil name.csv", csv_bytes, "text/csv")})
    assert up.status_code == 201
    upload = up.json()
    assert upload["filename"] == "evil_name.csv", "the path is stripped from uploaded names"

    preview = client.post(
        "/api/preview", json={"source": {"upload_id": upload["upload_id"]}}
    ).json()
    assert (preview["links"], preview["stores"]) == (2, 1)

    resp = client.post(
        "/api/runs",
        json={"source": {"upload_id": upload["upload_id"]}, "writers": ["json"]},
    )
    run = wait_until_finished(client, resp.json()["run_id"])
    assert run["source_kind"] == "file"
    assert run["source_name"] == "evil_name.csv"
    assert run["stores_done"] == 1


def test_unsupported_upload_is_refused(client):
    resp = client.post("/api/uploads", files={"file": ("links.pdf", b"x", "application/pdf")})
    assert resp.status_code == 415


def test_bad_input_and_settings_are_explained(client):
    assert (
        client.post("/api/preview", json={"source": {"text": "no links here"}}).json()["links"] == 0
    )
    resp = start(client, source={"text": "no links here"})
    assert resp.status_code == 422
    assert resp.json()["detail"]["code"] == "bad_input"
    resp = start(client, writers=["pdf"])
    assert resp.status_code == 422
    assert resp.json()["detail"]["code"] == "bad_settings"
    resp = client.post("/api/preview", json={"source": {"upload_id": "0" * 32}})
    assert resp.status_code == 404


@pytest.mark.parametrize("run_id", ["../../etc", "20261005T000000000Z-zzzzzz", "x"])
def test_unknown_or_malformed_run_ids_are_404(client, run_id):
    assert client.get(f"/api/runs/{run_id}").status_code == 404


def test_download_keys_are_a_fixed_list(client):
    run = wait_until_finished(client, start(client).json()["run_id"])
    assert client.get(f"/api/runs/{run['run_id']}/download/..%2Fconfig.json").status_code == 404
    assert client.get(f"/api/runs/{run['run_id']}/download/parquet").status_code == 404, (
        "not chosen"
    )


def test_events_stream_the_run_until_it_ends(client):
    run_id = start(client).json()["run_id"]
    with client.stream("GET", f"/api/runs/{run_id}/events") as resp:
        assert resp.headers["content-type"].startswith("text/event-stream")
        states = [
            json.loads(line.removeprefix("data: "))["state"]
            for line in resp.iter_lines()
            if line.startswith("data: ")
        ]
    assert states[-1] == "done"


def test_a_run_survives_a_server_restart(client, tmp_path):
    run = wait_until_finished(client, start(client).json()["run_id"])
    base = RunConfig.model_validate({"output": {"runs_dir": tmp_path / "runs"}})
    with TestClient(create_app(ApiSettings(base=base))) as fresh:
        again = fresh.get(f"/api/runs/{run['run_id']}").json()
    assert again["state"] == "done"
    assert again["stores_done"] == run["stores_done"]


def test_rows_preview_pages_through_a_table_without_raw_objects(client):
    run = wait_until_finished(client, start(client).json()["run_id"])
    body = client.get(f"/api/runs/{run['run_id']}/rows/products?limit=1").json()
    assert body["total"] == 1
    assert body["rows"][0]["title"] == "Cher Sweater"
    assert "raw" not in body["columns"], "the bulky source object is left out of the preview"
    stores = client.get(f"/api/runs/{run['run_id']}/rows/stores?offset=1&limit=1").json()
    assert (stores["total"], len(stores["rows"]), stores["offset"]) == (3, 1, 1)
    assert client.get(f"/api/runs/{run['run_id']}/rows/runs").status_code == 404


def test_rows_preview_shortens_long_text(client, tmp_path):
    long_page = "<html><body><p>" + "word " * 2000 + "</p></body></html>"
    settings = ApiSettings(
        base=RunConfig.model_validate({"output": {"runs_dir": tmp_path / "long"}}),
        poll_seconds=0.01,
    )
    app = create_app(
        settings, fetcher_factory=lambda: FakeFetcher({"https://long.com": (200, long_page)})
    )
    with TestClient(app) as c:
        run_id = c.post(
            "/api/runs", json={"source": {"text": "long.com"}, "writers": ["json"]}
        ).json()["run_id"]
        wait_until_finished(c, run_id)
        body = c.get(f"/api/runs/{run_id}/rows/pages").json()
    assert body["truncated_fields"] is True
    assert body["rows"][0]["text"].endswith("…")


def test_run_detail_carries_params_and_the_equivalent_cli(client):
    run = wait_until_finished(client, start(client).json()["run_id"])
    assert run["cli"] == "uv run scrapebot run links.txt -f xlsx,csv"
    assert run["config"]["input"]["text"].endswith("characters of pasted links"), (
        "pasted text is summarised"
    )


def test_overview_totals(client):
    wait_until_finished(client, start(client).json()["run_id"])
    body = client.get("/api/overview").json()
    assert (body["runs"], body["runs_done"], body["stores"], body["products"]) == (1, 1, 3, 1)
    assert body["last_run"]["mode"] == "full"


def test_stop_then_resume_from_the_interface(tmp_path):
    import threading

    gate = threading.Event()

    class GatedFetcher(FakeFetcher):
        def get(self, url):
            gate.wait(5)
            return super().get(url)

    base = RunConfig.model_validate(
        {
            "output": {"runs_dir": tmp_path / "runs"},
            "fetch": {"cache_dir": tmp_path / "c", "concurrency": 1},
        }
    )
    app = create_app(
        ApiSettings(base=base, poll_seconds=0.01), fetcher_factory=lambda: GatedFetcher(RESPONSES)
    )
    with TestClient(app) as c:
        body = {
            "source": {"text": LINKS},
            "writers": ["json"],
        }
        run_id = c.post("/api/runs", json=body).json()["run_id"]
        assert c.post(f"/api/runs/{run_id}/stop").status_code == 202
        gate.set()
        stopped = wait_until_finished(c, run_id)
        assert stopped["state"] == "stopped"
        assert stopped["stores_done"] < stopped["stores_total"]
        assert c.post(f"/api/runs/{run_id}/stop").status_code == 409

        assert c.post(f"/api/runs/{run_id}/resume").status_code == 202
        done = wait_until_finished(c, run_id)
        assert done["state"] == "done"
        assert done["stores_done"] == done["stores_total"] == 3
        assert c.post(f"/api/runs/{run_id}/resume").status_code == 409, (
            "a finished run cannot resume"
        )


SECRET = "sk-proj-THISISATESTKEYTHATMUSTNEVERLEAK123"


def test_llm_run_through_the_interface_never_leaks_the_key(tmp_path, caplog):
    """PRD section 13, absolute metric: zero API keys in logs or outputs."""
    import logging

    import respx

    from scrapebot.keys import RedactingFilter

    caplog.set_level(logging.DEBUG)
    caplog.handler.addFilter(RedactingFilter())
    grid = (
        "<html><body><h1>Rose</h1>"
        + "<p>Boutique.</p>" * 30
        + "<div>Cable Knit Cardigan $129.00</div></body></html>"
    )
    answer = {
        "store_type": "multi_brand",
        "products": [
            {"title": "Cable Knit Cardigan", "price": "$129.00", "source_url": "https://rose.com"}
        ],
        "confidence": 0.8,
    }
    completion = {
        "id": "x", "object": "chat.completion", "created": 1, "model": "gpt-4o-mini",
        "choices": [{"index": 0, "message": {"role": "assistant", "content": json.dumps(answer)}, "finish_reason": "stop"}],
        "usage": {"prompt_tokens": 500, "completion_tokens": 40, "total_tokens": 540},
    }  # fmt: skip
    base = RunConfig.model_validate(
        {"output": {"runs_dir": tmp_path / "runs"}, "fetch": {"cache_dir": tmp_path / "c"}}
    )
    settings = ApiSettings(base=base, poll_seconds=0.01, env_file=tmp_path / "missing.env")
    app = create_app(
        settings, fetcher_factory=lambda: FakeFetcher({"https://rose.com": (200, grid)})
    )
    with respx.mock(assert_all_called=False) as mock, TestClient(app) as c:
        route = mock.post("https://api.openai.com/v1/chat/completions").respond(
            200, json=completion
        )
        body = {
            "source": {"text": "rose.com"},
            "writers": ["json", "csv", "xlsx"],
            "llm": {
                "enabled": True,
                "model": "openai/gpt-4o-mini",
                "api_key": SECRET,
                "budget_usd": 1,
            },
        }
        resp = c.post("/api/runs", json=body)
        assert resp.status_code == 202
        run = wait_until_finished(c, resp.json()["run_id"])
        detail_text = c.get(f"/api/runs/{run['run_id']}").text

    assert route.called
    assert route.calls[0].request.headers["authorization"] == f"Bearer {SECRET}", (
        "the key reached the provider"
    )
    assert (run["llm_model"], run["llm_stores"], run["products"]) == ("openai/gpt-4o-mini", 1, 1)
    assert run["llm_cost_usd"] > 0
    assert "--no-llm" not in run["cli"], "the default model is on and needs no flag"

    leaks = [
        p
        for p in (tmp_path / "runs").rglob("*")
        if p.is_file() and SECRET.encode() in p.read_bytes()
    ]
    assert leaks == [], f"key written to {leaks}"
    assert SECRET not in detail_text
    assert SECRET not in caplog.text


def test_llm_switched_on_in_the_server_config_uses_the_env_key(tmp_path):
    """`scrapebot serve -c llm.yaml`: the interface sends no LLM settings, so the run
    takes them from the config, and the key from .env. It used to run with no key."""
    import respx

    answer = {
        "store_type": "multi_brand",
        "products": [
            {"title": "Cable Knit Cardigan", "price": "$129.00", "source_url": "https://rose.com"}
        ],
        "confidence": 0.8,
    }
    completion = {
        "id": "x", "object": "chat.completion", "created": 1, "model": "gpt-4o-mini",
        "choices": [{"index": 0, "message": {"role": "assistant", "content": json.dumps(answer)}, "finish_reason": "stop"}],
        "usage": {"prompt_tokens": 500, "completion_tokens": 40, "total_tokens": 540},
    }  # fmt: skip
    grid = (
        "<html><body>"
        + "<p>Boutique.</p>" * 30
        + "<div>Cable Knit Cardigan $129.00</div></body></html>"
    )
    env = tmp_path / ".env"
    env.write_text(f"OPENAI_API_KEY={SECRET}\n")
    base = RunConfig.model_validate(
        {
            "output": {"runs_dir": tmp_path / "runs"},
            "fetch": {"cache_dir": tmp_path / "c"},
            "llm": {"enabled": True, "model": "openai/gpt-4o-mini"},
        }
    )
    settings = ApiSettings(base=base, poll_seconds=0.01, env_file=env)
    app = create_app(
        settings, fetcher_factory=lambda: FakeFetcher({"https://rose.com": (200, grid)})
    )
    with respx.mock(assert_all_called=False) as mock, TestClient(app) as c:
        route = mock.post("https://api.openai.com/v1/chat/completions").respond(
            200, json=completion
        )
        body = {"source": {"text": "rose.com"}, "writers": ["json"]}
        run = wait_until_finished(c, c.post("/api/runs", json=body).json()["run_id"])

    assert route.called
    assert route.calls[0].request.headers["authorization"] == f"Bearer {SECRET}"
    assert run["products"] == 1


def test_llm_providers_report_env_keys_without_revealing_them(tmp_path):
    env = tmp_path / ".env"
    env.write_text(f"GEMINI_API_KEY={SECRET}\n")
    base = RunConfig.model_validate({"output": {"runs_dir": tmp_path / "runs"}})
    with TestClient(create_app(ApiSettings(base=base, env_file=env))) as c:
        resp = c.get("/api/llm/providers")
    providers = {p["provider"]: p for p in resp.json()}
    assert providers["gemini"]["key_in_env"] is True
    assert providers["openai"]["key_in_env"] is False
    assert providers["ollama"]["needs_key"] is False
    assert SECRET not in resp.text


def test_llm_check_explains_a_bad_key(tmp_path):
    import respx

    base = RunConfig.model_validate({"output": {"runs_dir": tmp_path / "runs"}})
    with (
        respx.mock(assert_all_called=False) as mock,
        TestClient(create_app(ApiSettings(base=base, env_file=tmp_path / "x"))) as c,
    ):
        mock.post("https://api.openai.com/v1/chat/completions").respond(
            401,
            json={
                "error": {
                    "message": "Incorrect API key provided: sk-proj-****",
                    "type": "invalid_request_error",
                    "param": None,
                    "code": "invalid_api_key",
                }
            },
        )
        resp = c.post(
            "/api/llm/check",
            json={"model": "openai/gpt-4o-mini", "api_key": "sk-proj-wrongwrongwrongwrong"},
        )
        bad_model = c.post("/api/llm/check", json={"model": "gpt-4o-mini"})
    assert resp.json() == {
        "ok": False,
        "message": "API key ditolak penyedia. Periksa key dan penyedianya.",
    }
    assert bad_model.json()["ok"] is False


def test_the_final_list_is_counted_previewed_and_downloadable(tmp_path):
    """ADR 0010: multi-brand stores that sell knitwear, beside the raw tables."""
    multi = json.dumps(
        {
            "products": [
                {"title": t, "vendor": v, "handle": f"p{i}", "variants": [{"price": "99.00"}]}
                for i, (t, v) in enumerate(
                    [("Aran Cardigan", "Vince"), ("Dress", "Ulla Johnson"), ("Belt", "Frame")]
                )
            ]
        }
    )
    responses = {
        **RESPONSES,
        "https://multi.com": (200, SHOPIFY_HOME),
        "https://multi.com/products.json?limit=250&page=1": (200, multi),
    }
    base = RunConfig.model_validate(
        {"output": {"runs_dir": tmp_path / "runs"}, "fetch": {"cache_dir": tmp_path / "cache"}}
    )
    settings = ApiSettings(base=base, uploads_dir=tmp_path / "uploads", poll_seconds=0.01)
    with TestClient(create_app(settings, fetcher_factory=lambda: FakeFetcher(responses))) as c:
        payload = {"source": {"text": LINKS + " https://multi.com"}, "writers": ["xlsx", "csv"]}
        run = wait_until_finished(c, c.post("/api/runs", json=payload).json()["run_id"])
        rows = c.get(f"/api/runs/{run['run_id']}/rows/final_products").json()
        download = c.get(f"/api/runs/{run['run_id']}/download/final_xlsx")

    assert (run["final_stores"], run["final_products"], run["knit_products"]) == (1, 1, 2)
    keys = {d["key"]: d["filename"] for d in run["downloads"]}
    assert keys["final_xlsx"] == f"{run['run_id']}-final.xlsx"
    assert "final_csv" in keys
    assert keys["knit_xlsx"] == f"{run['run_id']}-knit.xlsx", "monkees.com's sweater counts too"
    assert [r["title"] for r in rows["rows"]] == ["Aran Cardigan"]
    assert download.status_code == 200
    sheets = load_workbook(io.BytesIO(download.content), read_only=True).sheetnames
    assert sheets == ["final_stores", "final_products"]
