"""Every writer round-trips the same tables (standards, section 5).

Each format is read back with its own library and decoded by column kind; the
rows must equal the canonical rows exactly. Excel is the one documented
exception: cells over 32,767 characters are truncated.
"""

import csv
import json
import sqlite3
from pathlib import Path
from typing import Any

import duckdb
import pyarrow.parquet as pq
import pytest
from openpyxl import load_workbook

from scrapebot.outputs import available_writers, export, get_writer
from scrapebot.outputs.base import Table
from scrapebot.outputs.excel import MAX_CELL_CHARS, TRUNCATED_MARKER
from scrapebot.tables import (
    TABLES,
    Column,
    ContactRow,
    InputRow,
    LLMCallRow,
    PageRow,
    ProductRow,
    RunRow,
    StoreRow,
    columns,
)

RUN = "20261005T000000000Z-abc123"


def sample_tables(long_text: str = "Hello") -> list[Table]:
    rows = {
        "runs": [
            RunRow(
                run_id=RUN,
                started_at="2026-10-05T00:00:00+00:00",
                scrapebot_version="2.0",
                config={"limit": 2, "output": {"writers": ["csv"]}},
                limit=2,
            )
        ],
        "inputs": [
            InputRow(
                run_id=RUN,
                input_id=1,
                raw="https://café-tricot.fr/p?a=1&b=2",
                url="https://café-tricot.fr/p?a=1&b=2",
                domain="café-tricot.fr",
                status="processed",
                is_deep_link=True,
                meta={
                    "store_name": 'Café "Tricot"',
                    "lat": 26.1,
                    "tags": ["a", "b"],
                    "empty": None,
                },
            ),
            InputRow(run_id=RUN, input_id=2, raw="", status="no_website"),
        ],
        "stores": [
            StoreRow(
                run_id=RUN,
                domain="café-tricot.fr",
                url="https://café-tricot.fr",
                status="ok",
                layers_tried=["homepage", "shopify_feed"],
                product_count=2,
                ssl_bypassed=True,
                input_ids=[1],
                fetched_at="2026-10-05T00:00:01+00:00",
            )
        ],
        "products": [
            ProductRow(
                run_id=RUN,
                domain="café-tricot.fr",
                title='=SUM(A1:A2) pull-over, "laine"\tdouble',
                price_raw="1.234,50 €",
                currency="EUR",
                tags=["hiver", "laine"],
                source="jsonld",
                raw={"offers": {"price": "1.234,50"}, "n": [1, 2.5, None, True]},
            ),
            ProductRow(
                run_id=RUN,
                domain="café-tricot.fr",
                title="日本の セーター\nline two",
                source="shopify_feed",
                confidence=0.75,
            ),
        ],
        "pages": [
            PageRow(
                run_id=RUN,
                domain="café-tricot.fr",
                url="https://café-tricot.fr",
                page_kind="home",
                http_status=200,
                text=long_text,
            )
        ],
        "contacts": [
            ContactRow(
                run_id=RUN,
                domain="café-tricot.fr",
                type="email",
                value="bonjour@café-tricot.fr",
                source_url="https://café-tricot.fr/contact",
            )
        ],
        "changes": [],
        "llm_calls": [
            LLMCallRow(
                run_id=RUN,
                domain="café-tricot.fr",
                model="gemini/gemini-2.5-flash",
                prompt_version="extract-v1",
                status="ok",
                input_tokens=1200,
                output_tokens=310,
                cost_usd=0.000412,
                duration_seconds=2.4,
            )
        ],
    }
    return [
        Table(
            name=name, columns=columns(model), rows=[r.model_dump(mode="json") for r in rows[name]]
        )
        for name, model in TABLES.items()
    ]


def decode(value: Any, column: Column) -> Any:
    """A value read back from any format, as the canonical JSON value."""
    if value is None or value == "":
        return "" if column.kind == "str" else None
    if column.kind == "str":
        return str(value)
    if column.kind == "int":
        return int(value)
    if column.kind == "float":
        return float(value)
    if column.kind == "bool":
        return value if isinstance(value, bool) else str(value).lower() in ("true", "1")
    return json.loads(value) if isinstance(value, str) else value


def decoded(table: Table, records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return [{c.name: decode(r.get(c.name), c) for c in table.columns} for r in records]


def read_back(fmt: str, dest: Path, table: Table) -> list[dict[str, Any]]:
    name = table.name
    if fmt == "json":
        return json.loads((dest / "json" / f"{name}.json").read_text())
    if fmt == "jsonl":
        lines = (dest / "jsonl" / f"{name}.jsonl").read_text().splitlines()
        return [json.loads(line) for line in lines]
    if fmt in ("csv", "tsv"):
        with (dest / fmt / f"{name}.{fmt}").open(newline="") as fh:
            reader = csv.DictReader(fh, delimiter="," if fmt == "csv" else "\t")
            return decoded(table, list(reader))
    if fmt == "xlsx":
        ws = load_workbook(dest / "tables.xlsx", read_only=True)[name]
        values = list(ws.iter_rows(values_only=True))
        header = [str(h) for h in values[0]]
        return decoded(table, [dict(zip(header, row, strict=True)) for row in values[1:]])
    if fmt == "parquet":
        return decoded(table, pq.read_table(dest / "parquet" / f"{name}.parquet").to_pylist())
    if fmt == "sqlite":
        con = sqlite3.connect(dest / "tables.sqlite")
        con.row_factory = sqlite3.Row
        rows = [dict(r) for r in con.execute(f'SELECT * FROM "{name}"')]
        con.close()
        return decoded(table, rows)
    if fmt == "duckdb":
        con = duckdb.connect(str(dest / "tables.duckdb"), read_only=True)
        cur = con.execute(f'SELECT * FROM "{name}"')
        names = [d[0] for d in cur.description]
        rows = [dict(zip(names, r, strict=True)) for r in cur.fetchall()]
        con.close()
        return decoded(table, rows)
    raise AssertionError(fmt)


def test_every_registered_writer_has_a_round_trip_reader():
    assert set(available_writers()) == {
        "json",
        "jsonl",
        "csv",
        "tsv",
        "xlsx",
        "parquet",
        "sqlite",
        "duckdb",
    }


@pytest.mark.parametrize("fmt", available_writers())
def test_writer_round_trips_every_table(fmt, tmp_path):
    tables = sample_tables()
    get_writer(fmt).write(tables, tmp_path)
    for table in tables:
        assert read_back(fmt, tmp_path, table) == table.rows, f"{fmt}: {table.name}"


@pytest.mark.parametrize("fmt", ["csv", "tsv", "xlsx", "parquet", "sqlite", "duckdb"])
def test_empty_tables_keep_their_columns(fmt, tmp_path):
    tables = sample_tables()
    changes = next(t for t in tables if t.name == "changes")
    get_writer(fmt).write(tables, tmp_path)
    if fmt in ("csv", "tsv"):
        header = (tmp_path / fmt / f"changes.{fmt}").read_text().splitlines()[0]
        names = header.split("," if fmt == "csv" else "\t")
    elif fmt == "xlsx":
        names = list(
            next(
                load_workbook(tmp_path / "tables.xlsx", read_only=True)["changes"].iter_rows(
                    values_only=True
                )
            )
        )
    elif fmt == "parquet":
        names = pq.read_schema(tmp_path / "parquet" / "changes.parquet").names
    elif fmt == "sqlite":
        con = sqlite3.connect(tmp_path / "tables.sqlite")
        names = [r[1] for r in con.execute('PRAGMA table_info("changes")')]
        con.close()
    else:
        con = duckdb.connect(str(tmp_path / "tables.duckdb"), read_only=True)
        names = [d[0] for d in con.execute('SELECT * FROM "changes"').description]
        con.close()
    assert names == [c.name for c in changes.columns]


def test_excel_truncates_oversized_cells_and_keeps_formulas_as_text(tmp_path):
    long_text = "x" * (MAX_CELL_CHARS + 500)
    tables = sample_tables(long_text=long_text)
    get_writer("xlsx").write(tables, tmp_path)
    wb = load_workbook(tmp_path / "tables.xlsx")

    def cell(sheet: str, column: str):
        ws = wb[sheet]
        header = [c.value for c in ws[1]]
        return ws.cell(row=2, column=header.index(column) + 1)

    page_text = cell("pages", "text").value
    assert isinstance(page_text, str)
    assert len(page_text) == MAX_CELL_CHARS
    assert page_text.endswith(TRUNCATED_MARKER)
    title = cell("products", "title")
    assert str(title.value).startswith("=SUM")
    assert title.data_type == "s", "a title starting with '=' must not become a formula"


def test_other_formats_keep_oversized_text_whole(tmp_path):
    long_text = "x" * (MAX_CELL_CHARS + 500)
    tables = sample_tables(long_text=long_text)
    get_writer("parquet").write(tables, tmp_path)
    pages = pq.read_table(tmp_path / "parquet" / "pages.parquet").to_pylist()
    assert pages[0]["text"] == long_text


def test_export_writes_every_chosen_format(tmp_path):
    written = export(sample_tables(), ["csv", "parquet", "sqlite"], tmp_path)
    assert set(written) == {"csv", "parquet", "sqlite"}
    assert all(p.exists() for paths in written.values() for p in paths)


def test_export_survives_one_failing_writer(tmp_path, monkeypatch):
    def boom(tables, dest):
        raise RuntimeError("disk full")

    monkeypatch.setattr(get_writer("csv"), "write", boom)
    written = export(sample_tables(), ["csv", "json"], tmp_path)
    assert written["csv"] == []
    assert written["json"]


def test_excel_splits_a_table_across_sheets_at_the_row_limit(tmp_path, monkeypatch):
    writer = get_writer("xlsx")
    monkeypatch.setattr(writer, "max_rows", 2)
    rows = [
        ContactRow(
            run_id=RUN, domain="a.com", type="email", value=f"{i}@a.com", source_url="u"
        ).model_dump(mode="json")
        for i in range(5)
    ]
    table = Table(name="contacts", columns=columns(ContactRow), rows=rows)
    writer.write([table], tmp_path)
    wb = load_workbook(tmp_path / "tables.xlsx", read_only=True)
    assert wb.sheetnames == ["contacts", "contacts_2", "contacts_3"]
    assert sum(len(list(ws.iter_rows())) - 1 for ws in wb.worksheets) == 5


@pytest.mark.parametrize("fmt", available_writers())
def test_writers_stream_rows_from_a_one_shot_iterator(fmt, tmp_path):
    """A run's tables are read from disk as they are written; nothing needs a list."""
    tables = [
        Table(name=t.name, columns=t.columns, rows=iter(list(t.rows))) for t in sample_tables()
    ]
    get_writer(fmt).write(tables, tmp_path)
    for table in sample_tables():
        assert read_back(fmt, tmp_path, table) == list(table.rows), f"{fmt}: {table.name}"


def test_parquet_and_duckdb_write_large_tables_in_batches(tmp_path, monkeypatch):
    from scrapebot.outputs import base

    monkeypatch.setattr(base, "BATCH_ROWS", 3)
    rows = [
        ContactRow(
            run_id=RUN, domain="a.com", type="email", value=f"{i}@a.com", source_url="u"
        ).model_dump(mode="json")
        for i in range(10)
    ]
    table = Table(name="contacts", columns=columns(ContactRow), rows=rows)
    for fmt in ("parquet", "duckdb", "sqlite"):
        get_writer(fmt).write([table], tmp_path)
        assert read_back(fmt, tmp_path, table) == rows, fmt
