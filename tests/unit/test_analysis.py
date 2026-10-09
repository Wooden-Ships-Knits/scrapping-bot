"""The wholesale analysis (ADR 0011), offline: a small run folder written by the test."""

import json

import pytest
from openpyxl import load_workbook

from scrapebot import cli
from scrapebot.analysis.attributes import attributes, flag, gender_of, materials_of, price_of
from scrapebot.analysis.build import Options, analyze
from scrapebot.analysis.business import business_type
from scrapebot.analysis.inputs import (
    Customer,
    Inputs,
    core_name,
    domain_core,
    load,
    load_accounts,
    load_brands,
    load_stockists,
)
from scrapebot.analysis.location import Geocoder, Location, from_meta, from_pages, miles
from scrapebot.analysis.write import write

# --- product attributes ------------------------------------------------------------


def test_materials_largest_share_first_and_merino_wool_is_one_fibre():
    assert materials_of("30% Nylon, 70% Merino Wool") == ["merino", "synthetic"]
    assert materials_of("Cashmere blend with silk") == ["cashmere", "silk"]
    assert materials_of("Nylon 20% cashmere 80%") == ["cashmere", "synthetic"]


def test_unusual_letters_do_not_break_materials():
    """A real product, 2026-10-09: a Turkish dotless i in VISCOSE."""
    assert materials_of("100% VıSCOSE, 30% Cashmere") == ["cashmere"]


def test_gender_and_baby_alpaca_is_a_fibre_not_a_baby():
    assert gender_of("Women's Cable Sweater") == "women"
    assert gender_of("Men's Crew") == "men"
    assert gender_of("Kids Fair Isle Sweater") == "kids"
    assert gender_of("Baby Alpaca Scarf") == ""
    assert gender_of("Women's and Men's sizes") == "unisex"
    assert gender_of("Sweater", declared="female") == "women"


def test_shopify_sale_stock_and_launch_come_from_the_variants():
    a = attributes(
        {
            "title": "Aran Cardigan",
            "price_raw": "98.00",
            "source": "shopify_feed",
            "raw": {
                "published_at": "2026-09-01T10:00:00-04:00",
                "variants": [
                    {"available": False, "compare_at_price": "140.00"},
                    {"available": True, "compare_at_price": None},
                ],
            },
        }
    )
    assert (a.category, a.price, a.compare_at, a.on_sale, a.in_stock, a.launched) == (
        "cardigan", 98.0, 140.0, True, True, "2026-09-01",
    )  # fmt: skip


def test_prices_in_minor_units_are_read_as_written():
    assert price_of({"price_raw": "4800", "price_minor_unit": 2}) == 48.0
    woo = {"source": "woocommerce_feed", "price_raw": "3398",
           "raw": {"prices": {"currency_minor_unit": 2}}}  # fmt: skip
    assert price_of(woo) == 33.98, "older runs: WooCommerce says its minor unit in raw"
    bare = {"source": "render_json", "price_raw": "4800"}
    assert price_of(bare) is None, "cents or dollars cannot be told: left unknown"
    assert price_of({"source": "shopify_feed", "price_raw": "1,250.00"}) == 1250.0


def test_flags_written_as_text():
    assert (flag("true"), flag("0"), flag("InStock"), flag(""), flag(None)) == (
        True, False, True, None, None,
    )  # fmt: skip


# --- business type -----------------------------------------------------------------


@pytest.mark.parametrize(
    ("domain", "name", "text", "store_type", "kind"),
    [
        ("nordstrom.com", "", "", "multi_brand", "department_store"),
        ("garmentory.com", "", "", "multi_brand", "marketplace_resale"),
        ("betsyjenney.com", "", "", "multi_brand", "boutique"),  # not "etsy"
        ("keenelandshop.com", "", "", "multi_brand", "resort_hotel_club"),
        ("museumstore.org", "Museum Store", "", "multi_brand", "gift_museum"),
        ("currentboutique.com", "", "A consignment boutique", "multi_brand", "marketplace_resale"),
        ("lylasclothing.com", "", "Next to the golf shop on Main", "multi_brand", "boutique"),
        ("aran.com", "", "", "own_brand", "own_label"),
        ("15percentpledge.org", "", "", "", "non_retail"),
        ("thecashmeresale.com", "", "", "multi_brand", "boutique"),  # a sale, not resale
        ("countryclubprep.com", "", "", "multi_brand", "boutique"),  # a store's name
        ("whiteelephanthotel.com", "", "", "multi_brand", "resort_hotel_club"),
    ],
)
def test_business_types(domain, name, text, store_type, kind):
    products = 0 if kind == "non_retail" else 10
    assert business_type(domain, name, text, store_type, "ok", products).kind == kind


def test_page_phrases_are_hints_not_decisions():
    kind = business_type("lylasclothing.com", "", "Next to the golf shop", "multi_brand", "ok", 9)
    assert kind.kind == "boutique"
    assert kind.hints == ("resort_hotel_club: 'golf shop'",)


# --- location ------------------------------------------------------------------------


def test_address_from_the_contact_page_first():
    pages = [
        ("home", "Ship anywhere. Warehouse: Dallas, TX 75201"),
        ("contact", "Visit 345 Plandome Rd, Manhasset, NY 11030 or call"),
    ]
    found = from_pages(pages)
    assert found is not None
    assert (found.address, found.city, found.state, found.postal) == (
        "345 Plandome Rd", "Manhasset", "NY", "11030",
    )  # fmt: skip


def test_city_without_the_street_before_it_and_canadian_postcodes():
    found = from_pages([("contact", "123 Water Street Eau Claire, WI 54703")])
    assert found is not None
    assert found.city == "Eau Claire"
    canadian = from_pages([("about", "Wonderland Rd. S London, ON N6K 1L5")])
    assert canadian is not None
    assert (canadian.city, canadian.state, canadian.country) == ("London", "ON", "Canada")


def test_location_from_a_discovery_input_row():
    found = from_meta(
        {"name": "Rose", "city": "Naples", "state": "Florida", "postal_code": "34102"}
    )
    assert found is not None
    assert (found.state, found.postal, found.source) == ("FL", "34102", "input")
    assert from_meta({"name": "No place"}) is None


def test_geocoding_is_cached(tmp_path):
    calls = []
    geo = Geocoder(tmp_path / "geo.json", lookup=lambda q: calls.append(q) or (40.7, -74.0))
    place = Location(postal="10001", state="NY", country="United States")
    assert geo.point(place) == (40.7, -74.0)
    again = Geocoder(tmp_path / "geo.json", lookup=lambda q: calls.append(q) or None)
    assert again.point(place) == (40.7, -74.0)
    assert calls == ["10001, NY, United States"], "the second geocoder reads the cache"
    assert round(miles(40.7, -74.0, 40.8, -74.0), 1) == 6.9


# --- inputs ----------------------------------------------------------------------


def test_stockists_ignore_listing_sites_as_websites(tmp_path):
    path = tmp_path / "stockists.json"
    path.write_text(json.dumps([
        {"name": "Sunny Days", "website": "https://www.mapquest.com/us/x", "lat": 1, "lng": 2},
        {"name": "Hombadi Boutique", "website": "hombadi.com", "phone": "(401) 555-0100"},
    ]))  # fmt: skip
    sunny, hombadi = load_stockists(path)
    assert (sunny.domain, hombadi.domain, hombadi.phone) == ("", "hombadi.com", "4015550100")


def test_a_salesforce_account_export_as_it_comes(tmp_path):
    path = tmp_path / "accounts.csv"
    path.write_text(
        "Account Name,Website,Billing City,Billing State/Province,Billing Zip/Postal Code,"
        "Account Owner,Territory,Type\n"
        "Rose Boutique,www.roseboutique.com,Naples,FL,34102,Jane Rep,Southeast,Customer\n",
        encoding="utf-8",
    )
    [rose] = load_accounts(path)
    assert (rose.domain, rose.state, rose.postal, rose.rep, rose.territory, rose.status) == (
        "roseboutique.com", "FL", "34102", "Jane Rep", "Southeast", "Customer",
    )  # fmt: skip


def test_brands_with_relations_or_a_plain_list(tmp_path):
    table = tmp_path / "brands.csv"
    table.write_text("brand,relation\nKinross,peer\nAutumn Cashmere,peer\nCalLahan,competitor\n")
    assert load_brands(table) == ({"kinross", "autumncashmere"}, {"callahan"})
    plain = tmp_path / "brands.txt"
    plain.write_text("Kinross\nAlashan\n")
    assert load_brands(plain) == ({"kinross", "alashan"}, set())


def test_names_without_the_words_stores_add():
    assert core_name("The Tango Boutique") == "tango"
    assert domain_core("shoptangoboutique.com") == "tango"
    assert domain_core("hombadi.com") == "hombadi"


# --- the whole analysis ------------------------------------------------------------

RUN = "20261009T000000000Z-abcdef"


def product(domain, title, vendor, price="150.00", knit="garment", source="shopify_feed"):
    return {"run_id": RUN, "domain": domain, "title": title, "vendor": vendor,
            "price_raw": price, "currency": "USD", "product_type": "", "tags": [],
            "description": "", "url": "", "source": source, "knit_kind": knit,
            "is_knitwear": bool(knit), "raw": {}}  # fmt: skip


def store(domain, status="ok", store_type="multi_brand", products=0, input_ids=(0,)):
    return {"run_id": RUN, "domain": domain, "url": f"https://{domain}", "status": status,
            "store_type": store_type, "product_count": products, "currency": "USD",
            "platform": "shopify", "input_ids": list(input_ids)}  # fmt: skip


def write_run(tmp_path):
    root = tmp_path / RUN
    (root / "tables").mkdir(parents=True)
    tables = {
        "stores": [
            store("boutique.com", products=4, input_ids=[0]),
            store("near.com", products=1, input_ids=[1]),
            store("label.com", store_type="own_brand", products=6, input_ids=[2]),
            store("hombadi.com", products=1, input_ids=[3]),
            store("nordstrom.com", products=1, input_ids=[4]),
            store("fort.com", status="blocked", input_ids=[5]),
        ],
        "products": [
            product("boutique.com", "Aran Cardigan", "Kinross"),
            product("boutique.com", "Merino Crew", "Vince"),
            product("boutique.com", "Linen Dress", "Ulla Johnson", knit=""),
            product("boutique.com", "Silk Top", "Frame", knit=""),
            product("near.com", "Cable Sweater", "Vince"),
            *[product("label.com", f"Label Sweater {i}", "Label") for i in range(6)],
            product("hombadi.com", "Hombadi Sweater", "Wooden Ships"),
            product("nordstrom.com", "Cashmere Crew", "Nordstrom"),
        ],
        "contacts": [
            {
                "run_id": RUN,
                "domain": "boutique.com",
                "type": "email",
                "value": "hi@boutique.com",
                "source_url": "",
            }
        ],
        "pages": [
            {
                "run_id": RUN,
                "domain": "boutique.com",
                "url": "https://boutique.com/contact",
                "page_kind": "contact",
                "text": "1 Main St, Naples, FL 34102",
                "error": "",
            },
            {
                "run_id": RUN,
                "domain": "near.com",
                "url": "https://near.com/contact",
                "page_kind": "contact",
                "text": "Shop at 2 Bay St, Naples, FL 34103",
                "error": "",
            },
        ],
        "inputs": [{"run_id": RUN, "input_id": i, "meta": {}} for i in range(6)],
    }
    tables["inputs"][0]["meta"] = {"name": "Rose Boutique"}
    # near.com's name comes from its own markup, as the site wrote it
    tables["stores"][1]["identity"] = {"site_name": "Hill&#39;s Dry Goods"}
    for name, rows in tables.items():
        (root / "tables" / f"{name}.jsonl").write_text(
            "".join(json.dumps(r) + "\n" for r in rows), encoding="utf-8"
        )
    return root


POINTS = {"34102, FL, United States": (26.14, -81.79), "34103, FL, United States": (26.19, -81.80)}


def inputs():
    return Inputs(
        customers=[
            Customer("stockist", "Gulf Knits", city="Naples", state="FL", lat=26.19, lng=-81.80),
            Customer("stockist", "Far Away", state="WA", lat=47.6, lng=-122.3),
        ],
        peer_brands={"kinross"},
        competitor_brands={"callahan"},
        prices={"sweater": (75.0, 160.0), "cardigan": (80.0, 150.0)},
        sources={"stockists": "test", "brands": "test", "prices": "test"},
    )


def test_segments_scores_and_reasons(tmp_path):
    root = write_run(tmp_path)
    geocoder = Geocoder(tmp_path / "geo.json", lookup=POINTS.get)
    result = analyze(root, inputs(), Options(territory_miles=15, geocoder=geocoder))
    rows = {r.domain: r for r in result.stores}

    assert rows["fort.com"].segment == "not_read"
    assert rows["nordstrom.com"].segment == "not_a_fit"
    assert (rows["hombadi.com"].segment, rows["hombadi.com"].relationship) == (
        "existing_customer", "carries_wooden_ships",
    )  # fmt: skip
    assert rows["label.com"].segment == "competitor"
    boutique = rows["boutique.com"]
    assert (boutique.segment, boutique.store_name, boutique.city) == (
        "retail_partner", "Rose Boutique", "Naples",
    )  # fmt: skip
    assert boutique.peer_brands == ["Kinross"]
    assert (boutique.price_fit, boutique.knit_products) == ("fits", 2)
    assert boutique.nearest_stockist == "Gulf Knits"
    assert "-20 stockist Gulf Knits" in " ".join(boutique.reasons), "a stockist 3 miles away"
    assert "+10 carries peer brands: Kinross" in boutique.reasons
    assert boutique.score == 25 * 2 // 50 + 10 + 20 - 20 + 5 + 5  # knit, peers, price, territory...

    assert rows["near.com"].store_name == "Hill's Dry Goods"
    [competitor] = result.competitors
    assert (competitor.domain, competitor.label, competitor.knit_products) == (
        "label.com",
        "Label",
        6,
    )
    brands = {b.brand: b for b in result.brands}
    assert (brands["Kinross"].relation, brands["Wooden Ships"].relation) == ("peer", "wooden_ships")
    assert [k.title for k in result.knit if k.domain == "boutique.com"] == [
        "Aran Cardigan", "Merino Crew",
    ]  # fmt: skip


def test_missing_inputs_are_named_and_the_workbook_is_written(tmp_path):
    root = write_run(tmp_path)
    result = analyze(root, Inputs(), Options(geocoder=None))
    assert any(n.startswith("MISSING accounts") for n in result.notes)
    write(result, root, ["xlsx", "csv"])
    sheets = load_workbook(root / "analysis" / "tables.xlsx", read_only=True).sheetnames
    assert sheets == ["stores", "knit_products", "competitors", "brands"]
    assert "## Segments" in (root / "analysis" / "README.md").read_text()


def test_the_analyze_command(tmp_path, monkeypatch, capsys):
    root = write_run(tmp_path)
    folder = tmp_path / "inputs"
    folder.mkdir()
    (folder / "brands.csv").write_text("brand,relation\nKinross,peer\n")
    monkeypatch.chdir(tmp_path)
    assert cli.main(["analyze", str(root), "--inputs", str(folder), "--no-geocode"]) == 0
    out = capsys.readouterr().out
    assert "brands: " in out
    assert (root / "analysis" / "csv" / "stores.csv").exists()
    assert load(folder=folder).peer_brands == {"kinross"}
